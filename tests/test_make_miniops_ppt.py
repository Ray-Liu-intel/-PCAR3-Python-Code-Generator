import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation


REPO_ROOT = Path(__file__).resolve().parents[1]
GENERATOR_DIR = REPO_ROOT / "presentations" / "miniops"
sys.path.insert(0, str(GENERATOR_DIR))

import make_miniops_ppt as miniops  # noqa: E402


class MiniOpsPresentationTests(unittest.TestCase):
    def test_generate_english_deck_to_temp_dir(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)
            paths = list(miniops.generate("en", output_dir))

            self.assertEqual(1, len(paths))
            deck_path = paths[0]
            self.assertTrue(deck_path.exists())
            self.assertEqual("MiniOps_Management_EN.pptx", deck_path.name)

            prs = Presentation(deck_path)
            self.assertEqual(6, len(prs.slides))
            self.assertEqual(miniops.AUTHOR, prs.core_properties.author)
            self.assertEqual(miniops.EN_CONTENT["deck_title"], prs.core_properties.title)

    def test_generate_chinese_deck_to_temp_dir(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)
            paths = list(miniops.generate("cn", output_dir))

            self.assertEqual(1, len(paths))
            deck_path = paths[0]
            self.assertEqual("MiniOps_Management_CN.pptx", deck_path.name)

            prs = Presentation(deck_path)
            self.assertEqual(6, len(prs.slides))
            self.assertEqual(miniops.AUTHOR, prs.core_properties.author)
            self.assertEqual(miniops.CN_CONTENT["deck_title"], prs.core_properties.title)

    def test_generate_all_writes_both_decks(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)
            paths = sorted(path.name for path in miniops.generate("all", output_dir))
            self.assertEqual(
                ["MiniOps_Management_CN.pptx", "MiniOps_Management_EN.pptx"],
                paths,
            )

    def test_validate_deck_rejects_wrong_title_expectation(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_dir = Path(tmpdir)
            deck_path = next(miniops.generate("en", output_dir))

            with self.assertRaises(AssertionError):
                miniops.validate_deck(
                    deck_path,
                    expected_filename=deck_path.name,
                    expected_title="Wrong Title",
                )


if __name__ == "__main__":
    unittest.main()
