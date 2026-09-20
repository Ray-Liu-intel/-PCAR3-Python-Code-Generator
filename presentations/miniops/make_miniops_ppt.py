from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent
AUTHOR = "Ray Liu"
DATE_LABEL = "September 2026"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(9, 37, 64)
BLUE = RGBColor(22, 102, 181)
LIGHT_BLUE = RGBColor(228, 238, 248)
MID_BLUE = RGBColor(192, 215, 238)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(30, 41, 59)
GRAY = RGBColor(90, 101, 115)
BANNER = RGBColor(234, 244, 255)

CN_FONT_CANDIDATES = (
    "Noto Sans CJK SC",
    "Microsoft YaHei",
    "Source Han Sans SC",
    "SimSun",
    "WenQuanYi Micro Hei",
    "Arial Unicode MS",
    "Liberation Sans",
    "DejaVu Sans",
)

EN_FONT_CANDIDATES = (
    "Aptos",
    "Arial",
    "Liberation Sans",
    "DejaVu Sans",
    "Calibri",
)


@dataclass(frozen=True)
class Card:
    title: str
    bullets: tuple[str, ...]


@lru_cache(maxsize=1)
def available_font_families() -> set[str]:
    if not shutil.which("fc-list"):
        return set()
    try:
        output = subprocess.check_output(["fc-list", ":", "family"], text=True, errors="ignore")
    except subprocess.CalledProcessError:
        return set()
    return {
        family.strip().lower()
        for line in output.splitlines()
        for family in line.split(",")
        if family.strip()
    }


def resolve_font(deck_label: str, candidates: Sequence[str]) -> str:
    available = available_font_families()
    if available:
        for candidate in candidates:
            if candidate.lower() in available:
                if candidate != candidates[0]:
                    print(
                        f"[font] {deck_label}: preferred font '{candidates[0]}' not installed; using '{candidate}' instead.",
                        file=sys.stderr,
                    )
                return candidate
        raise RuntimeError(
            f"{deck_label}: none of the preferred fonts are installed: {', '.join(candidates)}"
        )
    print(
        f"[font] {deck_label}: fontconfig not available; using declared font '{candidates[0]}' without local verification.",
        file=sys.stderr,
    )
    return candidates[0]


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    lines: Sequence[str],
    *,
    font_name: str,
    font_size: float,
    color: RGBColor = DARK,
    bold_first: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    margin_left: float = 0.10,
    margin_right: float = 0.08,
    margin_top: float = 0.04,
    margin_bottom: float = 0.02,
    bullet: bool = False,
    line_spacing: float = 1.08,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Inches(margin_left)
    frame.margin_right = Inches(margin_right)
    frame.margin_top = Inches(margin_top)
    frame.margin_bottom = Inches(margin_bottom)
    for idx, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(1.5)
        paragraph.line_spacing = line_spacing
        if bullet:
            paragraph.level = 0
            paragraph.bullet = True
        run = paragraph.runs[0]
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.bold = bold_first and idx == 0
        run.font.color.rgb = color
    return box


def add_slide_base(slide, title: str, subtitle: str, *, font_name: str, slide_number: int):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(245, 248, 252)

    top_band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.82))
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = NAVY
    top_band.line.fill.background()

    accent_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(0.82), SLIDE_W, Inches(0.08)
    )
    accent_band.fill.solid()
    accent_band.fill.fore_color.rgb = BLUE
    accent_band.line.fill.background()

    add_textbox(
        slide,
        Inches(0.55),
        Inches(0.14),
        Inches(7.9),
        Inches(0.35),
        [title],
        font_name=font_name,
        font_size=24,
        color=WHITE,
        bold_first=True,
    )
    add_textbox(
        slide,
        Inches(0.57),
        Inches(0.47),
        Inches(8.1),
        Inches(0.20),
        [subtitle],
        font_name=font_name,
        font_size=10.5,
        color=RGBColor(220, 233, 247),
    )

    footer_line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(7.02), Inches(12.2), Inches(0.02)
    )
    footer_line.fill.solid()
    footer_line.fill.fore_color.rgb = MID_BLUE
    footer_line.line.fill.background()

    add_textbox(
        slide,
        Inches(0.55),
        Inches(7.05),
        Inches(4.2),
        Inches(0.18),
        [f"{AUTHOR} · {DATE_LABEL}"],
        font_name=font_name,
        font_size=9,
        color=GRAY,
    )
    add_textbox(
        slide,
        Inches(11.4),
        Inches(7.05),
        Inches(1.1),
        Inches(0.18),
        [str(slide_number)],
        font_name=font_name,
        font_size=9,
        color=GRAY,
        align=PP_ALIGN.RIGHT,
    )


def add_banner(slide, text: str, *, font_name: str, top: float):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(top), Inches(12.0), Inches(0.62)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BANNER
    shape.line.color.rgb = MID_BLUE
    shape.line.width = Pt(1)
    add_textbox(
        slide,
        Inches(0.82),
        Inches(top + 0.09),
        Inches(11.65),
        Inches(0.42),
        [text],
        font_name=font_name,
        font_size=14,
        color=NAVY,
        bold_first=True,
        align=PP_ALIGN.CENTER,
    )


def add_card(slide, left, top, width, height, card: Card, *, font_name: str, title_size: float, body_size: float):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = MID_BLUE
    shape.line.width = Pt(1.2)

    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()

    add_textbox(
        slide,
        left + Inches(0.08),
        top + Inches(0.12),
        width - Inches(0.16),
        Inches(0.36),
        [card.title],
        font_name=font_name,
        font_size=title_size,
        color=NAVY,
        bold_first=True,
    )
    add_textbox(
        slide,
        left + Inches(0.10),
        top + Inches(0.50),
        width - Inches(0.20),
        height - Inches(0.58),
        list(card.bullets),
        font_name=font_name,
        font_size=body_size,
        color=DARK,
        bullet=True,
        line_spacing=1.03,
    )


def set_notes(slide, notes: str):
    frame = slide.notes_slide.notes_text_frame
    frame.clear()
    lines = notes.splitlines()
    if not lines:
        return
    frame.text = lines[0]
    for line in lines[1:]:
        frame.add_paragraph().text = line


def build_summary_slide(prs: Presentation, content: dict, *, font_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_base(slide, content["title"], content["subtitle"], font_name=font_name, slide_number=1)
    add_banner(slide, content["headline"], font_name=font_name, top=1.05)

    positions = [
        (Inches(0.72), Inches(1.95)),
        (Inches(6.83), Inches(1.95)),
        (Inches(0.72), Inches(4.20)),
        (Inches(6.83), Inches(4.20)),
    ]
    for card, (left, top) in zip(content["cards"], positions):
        add_card(slide, left, top, Inches(5.78), Inches(1.95), card, font_name=font_name, title_size=14.5, body_size=11.5)
    set_notes(slide, content["notes"])


def build_problem_slide(prs: Presentation, content: dict, *, font_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_base(slide, content["title"], content["subtitle"], font_name=font_name, slide_number=2)

    positions = [Inches(0.72), Inches(4.37), Inches(8.02)]
    for card, left in zip(content["cards"], positions):
        add_card(slide, left, Inches(1.55), Inches(3.02), Inches(4.72), card, font_name=font_name, title_size=14.0, body_size=11.2)

    add_banner(slide, content["footer"], font_name=font_name, top=6.30)
    set_notes(slide, content["notes"])


def build_deliverables_slide(prs: Presentation, content: dict, *, font_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_base(slide, content["title"], content["subtitle"], font_name=font_name, slide_number=3)

    positions = [
        (Inches(0.72), Inches(1.35)),
        (Inches(6.83), Inches(1.35)),
        (Inches(0.72), Inches(4.13)),
        (Inches(6.83), Inches(4.13)),
    ]
    for card, (left, top) in zip(content["cards"], positions):
        add_card(slide, left, top, Inches(5.78), Inches(2.38), card, font_name=font_name, title_size=13.8, body_size=10.8)
    set_notes(slide, content["notes"])


def build_quality_slide(prs: Presentation, content: dict, *, font_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_base(slide, content["title"], content["subtitle"], font_name=font_name, slide_number=4)

    positions = [Inches(0.72), Inches(4.37), Inches(8.02)]
    for card, left in zip(content["cards"], positions):
        add_card(slide, left, Inches(1.60), Inches(3.02), Inches(4.45), card, font_name=font_name, title_size=13.8, body_size=10.9)

    boundary = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(6.12), Inches(11.85), Inches(0.72)
    )
    boundary.fill.solid()
    boundary.fill.fore_color.rgb = RGBColor(241, 246, 252)
    boundary.line.color.rgb = MID_BLUE
    add_textbox(
        slide,
        Inches(0.95),
        Inches(6.22),
        Inches(11.50),
        Inches(0.46),
        [content["footer"]],
        font_name=font_name,
        font_size=12.2,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )
    set_notes(slide, content["notes"])


def build_value_slide(prs: Presentation, content: dict, *, font_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_base(slide, content["title"], content["subtitle"], font_name=font_name, slide_number=5)

    rows = len(content["rows"]) + 1
    cols = 3
    table = slide.shapes.add_table(rows, cols, Inches(0.72), Inches(1.40), Inches(11.85), Inches(4.75)).table
    widths = [Inches(1.75), Inches(3.15), Inches(6.95)]
    for idx, width in enumerate(widths):
        table.columns[idx].width = width

    headers = content["headers"]
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.text = header
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.runs[0]
        run.font.name = font_name
        run.font.size = Pt(12.2)
        run.font.bold = True
        run.font.color.rgb = WHITE

    for row_idx, row in enumerate(content["rows"], start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_idx % 2 else LIGHT_BLUE
            frame = cell.text_frame
            frame.word_wrap = True
            frame.margin_left = Inches(0.08)
            frame.margin_right = Inches(0.05)
            frame.margin_top = Inches(0.05)
            frame.margin_bottom = Inches(0.03)
            paragraph = frame.paragraphs[0]
            paragraph.text = value
            paragraph.alignment = PP_ALIGN.LEFT
            run = paragraph.runs[0]
            run.font.name = font_name
            run.font.size = Pt(11.0)
            run.font.color.rgb = DARK

    add_banner(slide, content["footer"], font_name=font_name, top=6.32)
    set_notes(slide, content["notes"])


def build_next_steps_slide(prs: Presentation, content: dict, *, font_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_base(slide, content["title"], content["subtitle"], font_name=font_name, slide_number=6)

    left_card = Card(content["left_title"], tuple(content["left_bullets"]))
    right_card = Card(content["right_title"], tuple(content["right_bullets"]))
    add_card(slide, Inches(0.90), Inches(1.70), Inches(5.45), Inches(3.95), left_card, font_name=font_name, title_size=15.0, body_size=11.4)
    add_card(slide, Inches(7.00), Inches(1.70), Inches(5.45), Inches(3.95), right_card, font_name=font_name, title_size=15.0, body_size=11.4)
    add_banner(slide, content["footer"], font_name=font_name, top=6.05)
    set_notes(slide, content["notes"])


def configure_presentation(title: str) -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    core = prs.core_properties
    core.author = AUTHOR
    core.title = title
    core.subject = "MiniOps management presentation"
    core.comments = "Generated by make_miniops_ppt.py"
    core.category = "Presentation"
    return prs


def validate_deck(path: Path, *, expected_filename: str, expected_title: str):
    if path.name != expected_filename:
        raise AssertionError(f"expected filename {expected_filename}, got {path.name}")
    if path.suffix.lower() != ".pptx":
        raise AssertionError(f"{path.name}: expected .pptx suffix")
    if not path.exists() or path.stat().st_size <= 0:
        raise AssertionError(f"{path.name}: file was not written correctly")
    prs = Presentation(path)
    if prs.core_properties.author != AUTHOR:
        raise AssertionError(f"{path.name}: unexpected author {prs.core_properties.author!r}")
    if prs.core_properties.title != expected_title:
        raise AssertionError(f"{path.name}: unexpected title {prs.core_properties.title!r}")
    if len(prs.slides) != 6:
        raise AssertionError(f"{path.name}: expected 6 slides, found {len(prs.slides)}")
    title_found = False
    for shape in prs.slides[0].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip()
        if text and text.splitlines()[0] == expected_title:
            title_found = True
            break
    if not title_found:
        raise AssertionError(f"{path.name}: expected first-slide title {expected_title!r} not found")
    for index, slide in enumerate(prs.slides, start=1):
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if not notes:
            raise AssertionError(f"{path.name}: slide {index} notes are empty")

def write_deck(filename: str, font_name: str, content: dict, *, output_dir: Path):
    prs = configure_presentation(content["deck_title"])
    build_summary_slide(prs, content["slides"][0], font_name=font_name)
    build_problem_slide(prs, content["slides"][1], font_name=font_name)
    build_deliverables_slide(prs, content["slides"][2], font_name=font_name)
    build_quality_slide(prs, content["slides"][3], font_name=font_name)
    build_value_slide(prs, content["slides"][4], font_name=font_name)
    build_next_steps_slide(prs, content["slides"][5], font_name=font_name)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / filename
    prs.save(out_path)
    validate_deck(out_path, expected_filename=filename, expected_title=content["deck_title"])
    return out_path


CN_CONTENT = {
    "deck_title": "测试工程自动化：交付成果与业务价值",
    "slides": [
        {
            "title": "测试工程自动化：交付成果与业务价值",
            "subtitle": "MiniOps 管理层汇报｜10 分钟｜双语可编辑交付物",
            "headline": "四个工具已交付；下一步是验证团队级价值，而不是提前宣称收益数字。",
            "cards": (
                Card("可视化配置生成", ("PCAR3 Web UI 生成 Python 配置", "支持多 PLB、Loop Partition、条件结构")),
                Card("自然语言 / 构建辅助", ("Copilot Skill 生成 .pcar", "可选查询 Pattern、提交构建并检查真实 Pat 输出")),
                Card("DVS 预处理", ("CLI / Skill 处理 16 个 trigger 块", "支持 CCD / IOD / DRD，保留源文件")),
                Card("Scan Pin 检查", ("只读提取 ScanIn / ScanOut", "输出可追踪 JSON 报告用于人工复核")),
            ),
            "notes": """Timing: 1 minute
讲稿：这一页先给管理层一个结论。当前已经形成四个面向具体任务的工具：可视化配置生成、自然语言到 .pcar、DVS STIL 预处理，以及 scan pin 检查。这里可以明确区分两类结论：第一类是已经交付、可以演示和复现的工具能力；第二类是团队级效率、质量和采用价值，这些还需要后续用试点数据验证。今天不报告节省工时、效率百分比或 ROI，因为目前没有完成对等基线测量。

可确认内容：当前仓库中的 Web UI 功能已在本地读取代码核对；其余三项能力引用公开仓库的 README 或技能说明，不涉及内网材料。

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker""",
        },
        {
            "title": "解决的问题",
            "subtitle": "把重复规则工作转换为更一致、可复核的执行方式",
            "footer": "这些是任务特征，而不是已量化的历史损失；自动化支持工程判断，而不替代工程判断。",
            "cards": (
                Card("重复的配置组织", ("重复 PLB / Pattern / 条件配置", "对应方式：可视化或自然语言生成，减少重复结构录入")),
                Card("复杂且易错的规则", ("DVS cycle / trigger / pin 规则多", "对应方式：参数化处理和输入检查，避免静默偏差")),
                Card("重复的只读检查", ("反复做 scan-pin 提取和 group 对比", "对应方式：结构化检查结果和可追踪报告")),
            ),
            "notes": """Timing: 1 minute
讲稿：第二页解释为什么这些工作值得做，但仍然避免夸大。这里强调的是任务特征，例如 PLB 结构、Pattern 模式、DVS trigger 序列和 scan pin 归属检查本身具有重复性和规则性，所以适合工具化。这里不把这些任务直接描述成已经造成了多少损失，也不暗示自动化会替代工程判断。相反，更准确的说法是：工具先承担重复规则执行，工程师继续负责需求澄清、异常判断和最终接收。

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator/blob/main/README.md
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor/blob/main/README.md
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker/blob/main/README.md""",
        },
        {
            "title": "主要交付成果",
            "subtitle": "每项交付都区分“已交付能力”和“预期价值”",
            "cards": (
                Card("A. PCAR3 可视化 Web UI", ("已交付：多 PLB、Loop Partition、拖拽排序、状态持久化", "预期价值：减少手工结构整理与重复录入")),
                Card("B. 自然语言 Copilot Skill", ("已交付：生成 .pcar；可选解析 Pattern、提交构建、轮询状态并检查真实 Pat 输出", "边界：构建依赖合适环境与访问权限；仅代码生成不依赖构建环境")),
                Card("C. DVS CLI / Skill", ("已交付：转换预期 16 个 trigger 块，支持 CCD / IOD / DRD 与 .stil / .stil.gz", "预期价值：让重复编辑变成规则一致的文件处理")),
                Card("D. Scan Pin CLI / Skill", ("已交付：提取 ScanIn / ScanOut、标准化名称、检查 allowed group union、输出含链名/行号/匹配组的 JSON", "预期价值：提供可追踪的只读检查结果，便于人工复核")),
            ),
            "notes": """Timing: 3 minutes
讲稿：这一页是汇报核心，需要逐项讲清楚。第一项，当前仓库内的 Web UI 已从 HTML 代码核对，确实包含 loop partition、if/elif/else 条件块、拖拽和 localStorage 自动保存/恢复。第二项，自然语言 skill 的公开文档说明它可以生成 .pcar，并在具备 FLASH MCP 时继续做 pattern 解析、构建提交、状态轮询和实际 Pat 内容检查；这里要明确说明，这是公开仓库文档中描述的交付能力，不代表我在本次环境独立跑通了外部构建。第三和第四项同样要强调：交付的是规则、CLI 和技能能力；业务价值仍需通过试点测量。

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator/blob/main/Core%20Code2.0.HTML
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist/blob/main/README.md
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor/blob/main/README.md
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker/blob/main/README.md""",
        },
        {
            "title": "质量与风险控制",
            "subtitle": "强调输入保护、产物核对和边界说明，而不是过度承诺",
            "footer": "不宣称零缺陷或生产认证；最终接收与异常处置仍由人工负责。",
            "cards": (
                Card("保护输入", ("DVS 单独输出新文件，不覆盖源文件", "支持 dry-run，并拒绝异常、重复处理或会覆盖源文件的输入")),
                Card("核对产物", ("PList 关注真实 Pat 内容，而不仅是任务完成状态", "Scan pin 检查输出链名、行号和匹配组，便于追踪")),
                Card("明确边界", ("Pin PASS 只代表属于允许 group union", "不代表方向、时序、电气正确性或 tester readiness")),
            ),
            "notes": """Timing: 2 minutes
讲稿：这一页的重点是风险控制。管理层通常更关注边界是否说清楚。DVS 工具侧重保护输入和拒绝坏输入；PList 相关能力强调不能只看作业完成状态，而要检查是否真的有 Pat 内容；scan pin 检查强调结果是 membership check，而不是更广义的时序或电气签核。这里还要补充说明 scan pin 工具的四类状态：PASS 表示发现数据引脚且全部属于允许集合；FAIL 表示至少一个超出允许集合；UNDETERMINED 表示没有发现声明的数据引脚；ERROR 表示配置、解析或读取失败。这样可以防止管理层把 PASS 误解为“已可上机”。

Sources:
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist/blob/main/README.md
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor/blob/main/README.md
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker/blob/main/README.md""",
        },
        {
            "title": "业务价值与衡量方式",
            "subtitle": "先定义测量口径，再讨论是否具备推广价值",
            "headers": ("维度", "建议指标", "评估方式"),
            "rows": (
                ("效率", "同类任务的人工作业时间", "比较等价的人工/工具流程，包含准备、复核、纠错；机器等待单独记录"),
                ("质量", "返工与评审发现", "跟踪遗漏、误报、修正原因，而不只看是否生成成功"),
                ("采用", "实际使用频次 / 使用者", "区分作者本人使用与独立团队使用，再判断可推广性"),
                ("成本", "维护 / 培训 / 评审投入", "把维护和培训成本与节省一起评估，判断净价值"),
            ),
            "footer": "当前不宣称已测得的节省工时、改善百分比或 ROI。",
            "notes": """Timing: 2 minutes
讲稿：这页回答“结果怎么衡量”。为了避免价值叙述失真，建议把效率、质量、采用和成本拆开记录。效率不能只算脚本执行时间，而要用完整任务流程比较，包括准备、复核和纠错；同时把机器等待和人工操作分开。质量要记录返工原因、误报和遗漏，而不是只统计是否成功生成。采用要区分作者本人使用和独立团队使用。成本则要把维护、培训和评审投入算进去。只有这样，后续说“值得推广”才有依据。

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker""",
        },
        {
            "title": "下一步与所需支持",
            "subtitle": "建议的小范围试点方案：先验证，再扩展",
            "left_title": "建议试点步骤",
            "left_bullets": (
                "选择规则稳定、频率较高的试点任务，并指定试点负责人和验收负责人",
                "准备经过批准的合成或脱敏示例，避免使用专有客户数据",
                "建立人工/工具基线，以及异常、复核和返工记录",
                "按约定验收标准评估效率、质量和维护成本，再决定是否扩展",
            ),
            "right_title": "需要的管理支持",
            "right_bullets": (
                "确认试点参与人员和业务窗口",
                "批准可用于演示和验证的样例范围",
                "安排结果复盘，统一判断是否进入下一阶段",
                "在扩展前明确哪些结论需要独立团队复核",
            ),
            "footer": "这是建议方案，尚未批准；扩展应建立在试点结果和约定验收标准之上。",
            "notes": """Timing: 1 minute
讲稿：最后一页只提一个可执行的下一步：用小范围试点来回答价值问题。建议从规则稳定、发生频率高的任务开始，提前指定试点和验收负责人，并使用经过批准的合成或脱敏样例，避免任何客户或内部敏感数据外露。试点期间，把人工基线、工具结果、异常和返工原因都记录下来。完成后再做一次结果复盘，决定是否扩大范围。这样既避免一次性大规模推广的风险，也能给后续资源投入提供更清晰依据。

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker""",
        },
    ],
}


EN_CONTENT = {
    "deck_title": "Test Engineering Automation: Deliverables and Business Value",
    "slides": [
        {
            "title": "Test Engineering Automation: Deliverables and Business Value",
            "subtitle": "MiniOps management presentation — 10 minutes — editable bilingual deliverables",
            "headline": "Four tools are delivered; the next step is to validate team-level value rather than claim benefit numbers early.",
            "cards": (
                Card("Visual configuration generation", ("PCAR3 Web UI generates Python configuration", "Supports multiple PLBs, Loop Partition, and conditional structures")),
                Card("Natural language / build assistance", ("Copilot Skill generates .pcar", "Can optionally resolve patterns, submit builds, and verify real Pat output")),
                Card("DVS preprocessing", ("CLI / Skill expands 16 trigger blocks", "Supports CCD / IOD / DRD while preserving source files")),
                Card("Scan pin checking", ("Read-only extraction of ScanIn / ScanOut", "Produces traceable JSON output for human review")),
            ),
            "notes": """Timing: 1 minute
Talk track: Start with the management conclusion. We now have four task-specific tools: visual configuration generation, natural-language-to-.pcar assistance, DVS STIL preprocessing, and scan-pin checking. I want to separate two categories of statements very clearly. First, we have delivered tool capabilities that can be demonstrated and reproduced. Second, team-level efficiency, quality, and adoption outcomes still need pilot measurement before they should be claimed. So this presentation does not report hours saved, percentage improvement, or ROI. Those require equivalent-task baselines and review data.

What is independently checked here: the current repository's Web UI features were verified by reading the local HTML implementation. The other three tool descriptions are based on public GitHub repository documentation rather than private systems or internal material.

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker""",
        },
        {
            "title": "Problems Addressed",
            "subtitle": "Turn repeated rule-heavy work into more consistent and reviewable execution",
            "footer": "These are task characteristics, not quantified historical losses; automation supports engineering judgment rather than replacing it.",
            "cards": (
                Card("Repeated configuration work", ("Repeated PLB / pattern / conditional setup", "Approach: visual or natural-language generation to reduce repeated structuring")),
                Card("Complex rule execution", ("DVS cycle, trigger, and pin rules are detailed and easy to misapply", "Approach: parameterized processing plus input checks instead of ad hoc edits")),
                Card("Repeated read-only checking", ("Recurring scan-pin extraction and group comparison", "Approach: structured checks and traceable reporting for review")),
            ),
            "notes": """Timing: 1 minute
Talk track: This slide explains why these tasks are good candidates for automation without overstating the business case. The point is not that we already measured a historical loss number. The point is that PLB structures, pattern modes, DVS trigger sequences, and scan-pin membership checks all have repeated, rule-based characteristics. That makes them suitable for tools. The role of the engineer does not disappear. Engineers still own requirement clarification, exceptions, and final acceptance. The tools simply move repeated rule execution into a more consistent form.

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator/blob/main/README.md
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor/blob/main/README.md
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker/blob/main/README.md""",
        },
        {
            "title": "Key Deliverables",
            "subtitle": "Each deliverable separates shipped capability from intended value",
            "cards": (
                Card("A. PCAR3 visual Web UI", ("Delivered: multiple PLBs, Loop Partition, drag/drop ordering, and state persistence", "Intended value: less manual structuring and less repeated re-entry")),
                Card("B. Natural-language Copilot Skill", ("Delivered: generate .pcar, and optionally resolve patterns, submit builds, poll status, and verify real Pat output", "Boundary: builds require the right environment and access; code-only generation does not")),
                Card("C. DVS CLI / Skill", ("Delivered: transform the expected 16 trigger blocks, support CCD / IOD / DRD, and handle .stil / .stil.gz", "Intended value: convert repeated editing into consistent rule-based processing")),
                Card("D. Scan pin CLI / Skill", ("Delivered: extract ScanIn / ScanOut, normalize names, check allowed-group union, and emit JSON with chain names, line numbers, and matches", "Intended value: produce traceable read-only check results for human review")),
            ),
            "notes": """Timing: 3 minutes
Talk track: This is the core slide, so I would walk through each tool briefly but precisely. First, the current repository's Web UI was verified from the local HTML implementation and does contain loop partition, if/elif/else support, drag-and-drop behavior, and localStorage save/restore logic. Second, the natural-language skill's public documentation says it can generate .pcar and, when FLASH MCP is available, continue with pattern resolution, build submission, status polling, and real Pat-content verification. I am presenting that as a documented deliverable capability from the public repository, not as an external build independently executed in this environment. The DVS and scan-pin items follow the same pattern: shipped rules and tooling are the current deliverable; business outcome still requires pilot measurement.

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator/blob/main/Core%20Code2.0.HTML
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist/blob/main/README.md
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor/blob/main/README.md
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker/blob/main/README.md""",
        },
        {
            "title": "Quality and Risk Controls",
            "subtitle": "Emphasize input protection, artifact checks, and explicit boundaries rather than over-promising",
            "footer": "No claim of zero defects or production certification; final acceptance and exception handling remain human-owned.",
            "cards": (
                Card("Protect inputs", ("DVS writes a separate output instead of overwriting the source", "Supports dry-run and rejects malformed, reprocessed, or source-overwriting input")),
                Card("Verify artifacts", ("PList-oriented flow checks for real Pat content, not only a completed job status", "Scan-pin reports include chain names, line numbers, and matching groups for traceability")),
                Card("State boundaries explicitly", ("A pin PASS means membership in the allowed group union only", "It does not mean direction, timing, electrical correctness, or tester readiness")),
            ),
            "notes": """Timing: 2 minutes
Talk track: This slide is mainly about control and boundaries. Management usually wants to know whether the risks are framed clearly. The DVS tool focuses on input protection and explicit rejection of bad or already processed input. The PList-related flow emphasizes that completion status alone is not enough; the artifact still needs content verification. The scan-pin checker makes its scope explicit: it is a membership check, not a broader signoff for timing or tester execution. In the notes I also want to name the four result states so there is no ambiguity: PASS means data pins were found and all are in the allowed union; FAIL means at least one declared data pin is outside it; UNDETERMINED means no declared data pins were found; ERROR means configuration, parsing, or file-read failure.

Sources:
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist/blob/main/README.md
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor/blob/main/README.md
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker/blob/main/README.md""",
        },
        {
            "title": "Business Value and Measurement",
            "subtitle": "Define the measurement method first, then decide whether the tools justify broader rollout",
            "headers": ("Dimension", "Suggested metric", "How to evaluate"),
            "rows": (
                ("Efficiency", "Hands-on time for comparable tasks", "Compare equivalent manual and tool-assisted workflows, including preparation, review, and correction; record machine wait separately"),
                ("Quality", "Rework and review findings", "Track misses, false alarms, and correction causes rather than only whether generation succeeded"),
                ("Adoption", "Actual usage and users", "Separate author usage from independent team usage before claiming broader value"),
                ("Cost", "Maintenance, training, and review effort", "Assess support cost alongside any savings to judge net value")),
            "footer": "No measured hours saved, percentage improvement, or ROI is claimed yet.",
            "notes": """Timing: 2 minutes
Talk track: This slide answers the management question, how should value be measured. To avoid distorted conclusions, I recommend tracking efficiency, quality, adoption, and cost separately. Efficiency should be based on full task flow, not just script runtime, and should include preparation, review, and correction. Quality should capture misses, false alarms, and rework causes, not just successful output generation. Adoption should distinguish the author's own use from independent team use. Cost should include maintenance, training, and review effort. Only after those dimensions are visible should we decide whether the tools deserve a wider rollout.

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker""",
        },
        {
            "title": "Next Steps and Support Needed",
            "subtitle": "Proposed small-scope pilot plan: validate first, then expand",
            "left_title": "Proposed pilot steps",
            "left_bullets": (
                "Select stable-rule, high-frequency pilot tasks and name both pilot and acceptance owners",
                "Use approved synthetic or sanitized examples instead of proprietary customer data",
                "Establish a manual/tool baseline plus exception, review, and rework records",
                "Evaluate efficiency, quality, and maintenance against agreed acceptance criteria before expanding",
            ),
            "right_title": "Support requested",
            "right_bullets": (
                "Confirm pilot participants and business contacts",
                "Approve the example scope that can be used for demonstration and validation",
                "Schedule a results review to decide whether a next phase is justified",
                "Agree which conclusions require independent team validation before expansion",
            ),
            "footer": "This is a proposed plan, not an approved rollout; expansion should follow pilot results and agreed acceptance criteria.",
            "notes": """Timing: 1 minute
Talk track: The final ask is intentionally small and practical. Rather than trying to scale immediately, I recommend a limited pilot on stable-rule, high-frequency tasks with named pilot and acceptance owners. The examples should be approved synthetic or sanitized data only. During the pilot, capture the manual baseline, tool-assisted results, exceptions, and rework causes. Then hold a results review and decide whether a broader next phase is justified. That gives management a lower-risk way to validate value before further investment.

Sources:
- https://github.com/Ray-Liu-intel/-PCAR3-Python-Code-Generator
- https://github.com/Ray-Liu-intel/PCAR3-Natural-Language-to-Plist
- https://github.com/Ray-Liu-intel/DVS-STIL-Preprocessor
- https://github.com/Ray-Liu-intel/stil-scan-pin-checker""",
        },
    ],
}


def generate(language: str, output_dir: Path) -> Iterable[Path]:
    if language in {"cn", "all"}:
        cn_font = resolve_font("Chinese deck", CN_FONT_CANDIDATES)
        yield write_deck("MiniOps_Management_CN.pptx", cn_font, CN_CONTENT, output_dir=output_dir)
    if language in {"en", "all"}:
        en_font = resolve_font("English deck", EN_FONT_CANDIDATES)
        yield write_deck("MiniOps_Management_EN.pptx", en_font, EN_CONTENT, output_dir=output_dir)


def main():
    parser = argparse.ArgumentParser(description="Generate editable bilingual MiniOps management PowerPoint decks.")
    parser.add_argument(
        "--language",
        choices=("all", "cn", "en"),
        default="all",
        help="Select which deck to generate. Defaults to all.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for generated decks. Defaults to {DEFAULT_OUTPUT_DIR}.",
    )
    args = parser.parse_args()
    paths = list(generate(args.language, args.output_dir.resolve()))
    for path in paths:
        print(f"Generated and validated: {path}")


if __name__ == "__main__":
    main()
